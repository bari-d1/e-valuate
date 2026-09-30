"""Tests for upload text extraction (no DB)."""

import io

from docx import Document

from app.services.document_text import extract_docx_text, extract_text_from_upload


def test_extract_text_from_upload_txt():
    text = extract_text_from_upload("notes.txt", b"Board oversight of risk is essential.")
    assert "risk" in text


def test_extract_text_from_upload_docx():
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("The board receives timely information for effective oversight.")
    doc.add_table(rows=1, cols=1)
    doc.tables[0].cell(0, 0).text = "Committee oversight is adequate."
    doc.save(buf)

    text = extract_text_from_upload("instrument.docx", buf.getvalue())
    assert "timely information" in text
    assert "Committee oversight" in text


def test_extract_docx_text_paragraphs_only():
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Line one.")
    doc.add_paragraph("Line two.")
    doc.save(buf)

    assert extract_docx_text(buf.getvalue()) == "Line one.\nLine two."
