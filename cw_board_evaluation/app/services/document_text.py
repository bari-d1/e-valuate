"""
Extract plain text from uploaded documents.

- extract_text_from_upload: AI questionnaire generation uploads (.txt/.md/.docx; unchanged).
- extract_document_text: organisation documents (.pdf/.docx/.txt).
"""

from __future__ import annotations

import io

from fastapi import HTTPException


def _normalize_filename(filename: str) -> str:
    return (filename or "").strip().lower()


def extract_docx_text(raw: bytes) -> str:
    """Pull paragraph and table cell text from a .docx file."""
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts: list[str] = []

    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = (cell.text or "").strip()
                if text:
                    parts.append(text)

    return "\n".join(parts).strip()


def extract_text_from_upload(filename: str, raw: bytes) -> str:
    """
    Return UTF-8 plain text from supported upload types.

    Supported: .txt, .md, .markdown, .docx
    """
    if not raw:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    name = _normalize_filename(filename)

    if name.endswith((".txt", ".md", ".markdown")):
        return raw.decode("utf-8", errors="ignore").strip()

    if name.endswith(".docx"):
        try:
            text = extract_docx_text(raw)
        except Exception as exc:
            raise HTTPException(
                status_code=400,
                detail=f"Could not read DOCX file: {exc}",
            ) from exc
        if not text:
            raise HTTPException(
                status_code=400,
                detail="DOCX file contained no extractable text (empty document).",
            )
        return text

    if name.endswith(".doc"):
        raise HTTPException(
            status_code=400,
            detail="Legacy .doc files are not supported. Save as .docx or export to .txt and upload again.",
        )

    if name.endswith(".pdf"):
        raise HTTPException(
            status_code=400,
            detail="PDF extraction is not supported yet. Save as .docx or .txt and upload again.",
        )

    raise HTTPException(
        status_code=400,
        detail="Unsupported file type. Upload .txt, .md, or .docx.",
    )


class DocumentTextError(ValueError):
    """Text could not be extracted. The message is safe to show to users."""


def extract_pdf_text(raw: bytes) -> str:
    """Text from a PDF's text layer (pypdf). Scanned pages without a text layer yield nothing."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise DocumentTextError("The PDF is password-protected.") from exc
    pages = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(text)
    return "\n\n".join(pages).strip()


def extract_document_text(filename: str, raw: bytes) -> str:
    """
    Plain text from an organisation document (.pdf, .docx, .txt).
    Raises DocumentTextError when nothing usable can be read.
    """
    name = _normalize_filename(filename)
    try:
        if name.endswith(".pdf"):
            text = extract_pdf_text(raw)
        elif name.endswith(".docx"):
            text = extract_docx_text(raw)
        elif name.endswith(".txt"):
            text = raw.decode("utf-8", errors="replace").strip()
        else:
            raise DocumentTextError("Unsupported file type.")
    except DocumentTextError:
        raise
    except Exception as exc:
        raise DocumentTextError("The file could not be read.") from exc
    if not text:
        raise DocumentTextError("No text found in the file (it may be a scanned image).")
    return text
